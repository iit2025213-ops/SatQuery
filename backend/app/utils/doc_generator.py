import io
import re
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import logging

logger = logging.getLogger(__name__)

class DocGenerator:
    """Utility class to generate Microsoft Word (.docx) documents."""

    @staticmethod
    def generate_timeline_docx(ai_summary: str, timeline_data: dict, video_url: str = None, index_images: dict = None) -> bytes:
        """
        Generates a DOCX file containing the AI Markdown summary and the raw GEE timeline data.
        Can optionally embed raw image bytes for indices.
        Returns the raw bytes of the DOCX file.
        """
        document = Document()
        
        # Add a stylish main title
        title = document.add_heading('SatQuery AI - Temporal Analysis Report', 0)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        document.add_paragraph("Generated automatically by SatQuery Geospatial AI Pipeline.\n")
        
        # We will parse the Markdown AI Summary simply by looking at line prefixes
        lines = ai_summary.split("\n")
        for line in lines:
            line = line.strip()
            if not line:
                continue
                
            # Skip the markdown image embedding line for the video, as we add it explicitly later
            if line.startswith("![") and "Time-Lapse" in line:
                continue
            
            # Headings
            if line.startswith("# "):
                document.add_heading(line[2:].replace("**", ""), level=1)
            elif line.startswith("## "):
                document.add_heading(line[3:].replace("**", ""), level=2)
            elif line.startswith("### "):
                document.add_heading(line[4:].replace("**", ""), level=3)
            
            # Bullet points
            elif line.startswith("- ") or line.startswith("* "):
                p = document.add_paragraph(style='List Bullet')
                # basic bold parsing
                text = line[2:]
                DocGenerator._add_parsed_text(p, text)
            
            # Standard paragraph
            else:
                p = document.add_paragraph()
                DocGenerator._add_parsed_text(p, line)
                
        document.add_page_break()
        
        # Insert raw GEE stats as a table
        document.add_heading('Raw Geospatial Statistics (GEE)', level=1)
        frames = timeline_data.get("frames", [])
        if frames:
            table = document.add_table(rows=1, cols=5)
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Date'
            hdr_cells[1].text = 'NDVI (Veg)'
            hdr_cells[2].text = 'NDWI (Water)'
            hdr_cells[3].text = 'NDBI (Urban)'
            hdr_cells[4].text = 'NBR (Burn)'
            
            for frame in frames:
                row_cells = table.add_row().cells
                row_cells[0].text = str(frame.get("date", ""))
                row_cells[1].text = str(frame.get("ndvi_mean", "N/A"))
                row_cells[2].text = str(frame.get("ndwi_mean", "N/A"))
                row_cells[3].text = str(frame.get("ndbi_mean", "N/A"))
                row_cells[4].text = str(frame.get("nbr_mean", "N/A"))
                
        document.add_paragraph()
        
        # Define metadata and color legends for each index
        INDEX_METADATA = {
            "NDVI": {
                "name": "Normalized Difference Vegetation Index (NDVI)",
                "desc": "Measures the density and health of vegetation in the area.",
                "legend": "🔴 Red/Orange: Barren land or built-up areas | 🟢 Green: Dense, healthy vegetation."
            },
            "NDWI": {
                "name": "Normalized Difference Water Index (NDWI)",
                "desc": "Highlights water bodies and tracks moisture content in plants and soil.",
                "legend": "⚪ White/Light Blue: Dry land or shallow water | 🔵 Deep Blue: Deep, clear water bodies."
            },
            "NDBI": {
                "name": "Normalized Difference Built-up Index (NDBI)",
                "desc": "Emphasizes urban areas, buildings, and impervious surfaces.",
                "legend": "⚪ White/Yellow: Non-urban areas or vegetation | 🔴 Red: Dense urban built-up areas."
            },
            "NBR": {
                "name": "Normalized Burn Ratio (NBR)",
                "desc": "Used to identify burned areas, fire severity, and vegetation survival.",
                "legend": "🟡 Yellow/Orange: Healthy vegetation or unburned land | ⚫ Dark Red/Black: Severe burn scars."
            }
        }
        
        # Add index images if provided
        if index_images:
            document.add_page_break()
            document.add_heading('Latest Geospatial Index Visualizations', level=1)
            for idx_name, img_bytes in index_images.items():
                if img_bytes:
                    meta = INDEX_METADATA.get(idx_name, {"name": idx_name, "desc": "", "legend": ""})
                    document.add_heading(f'{meta["name"]}', level=2)
                    
                    p_desc = document.add_paragraph(meta["desc"])
                    p_desc.style = 'List Bullet'
                    
                    p_leg = document.add_paragraph(meta["legend"])
                    p_leg.style = 'List Bullet'
                    
                    img_stream = io.BytesIO(img_bytes)
                    try:
                        document.add_picture(img_stream, width=Inches(4.5))
                    except Exception as e:
                        logger.error(f"Failed to embed {idx_name} image: {e}")
                    document.add_paragraph()
        
        # Add the video URL reference
        if video_url:
            document.add_heading('Time-Lapse Video', level=1)
            p = document.add_paragraph()
            p.add_run("View the generated time-lapse animation here: ")
            p.add_run(video_url).underline = True
            
        # Save to memory
        docx_stream = io.BytesIO()
        document.save(docx_stream)
        docx_stream.seek(0)
        return docx_stream.read()

    @staticmethod
    def generate_spatial_docx(ai_summary: str, spatial_stats: dict, mapbox_img_bytes: bytes = None, index_images: dict = None) -> bytes:
        """
        Generates a DOCX file containing the AI Markdown summary and the raw GEE spatial data.
        Embeds a Mapbox image for visual reference and GEE-derived index thumbnails.
        """
        document = Document()
        
        title = document.add_heading('SatQuery AI - Spatial Analysis Report', 0)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        document.add_paragraph("Generated automatically by SatQuery Geospatial AI Pipeline.\n")
        
        if mapbox_img_bytes:
            document.add_heading('Mapbox Visual Reference (Visualization Only)', level=1)
            document.add_paragraph("This high-resolution satellite imagery is provided by Mapbox strictly for visual context. It is NOT used for any scientific calculations or index measurements.")
            img_stream = io.BytesIO(mapbox_img_bytes)
            try:
                document.add_picture(img_stream, width=Inches(5))
            except Exception as e:
                logger.error(f"Failed to embed Mapbox image: {e}")
            document.add_paragraph()
            
        # Parse the Markdown AI Summary
        lines = ai_summary.split("\n")
        
        # Determine if we should inject provenance
        provenance = kwargs.get("provenance") if "provenance" in kwargs else None
        if provenance and isinstance(provenance, dict):
            job_id = provenance.get("job_id")
            evidence_ids = provenance.get("evidence_ids", [])
            if evidence_ids:
                lines.append("")
                lines.append("## Evidence & Provenance Trail")
                lines.append(f"Job ID: {job_id}")
                for ev_id in evidence_ids:
                    lines.append(f"- GEE Scientific Evidence ID: {ev_id}")

        for line in lines:
            line = line.strip()
            if not line:
                continue
            
            if line.startswith("# "):
                document.add_heading(line[2:].replace("**", ""), level=1)
            elif line.startswith("## "):
                document.add_heading(line[3:].replace("**", ""), level=2)
            elif line.startswith("### "):
                document.add_heading(line[4:].replace("**", ""), level=3)
            elif line.startswith("- ") or line.startswith("* "):
                p = document.add_paragraph(style='List Bullet')
                DocGenerator._add_parsed_text(p, line[2:])
            else:
                p = document.add_paragraph()
                DocGenerator._add_parsed_text(p, line)
                
        document.add_page_break()
        
        document.add_heading('Raw Geospatial Statistics (GEE-Derived)', level=1)
        document.add_paragraph("All scientific metrics, indices, and findings are derived directly from Google Earth Engine satellite observations.")
        
        table = document.add_table(rows=2, cols=4)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'NDVI (Vegetation)'
        hdr_cells[1].text = 'NDWI (Water)'
        hdr_cells[2].text = 'NDBI (Urban)'
        hdr_cells[3].text = 'NBR (Burn)'
        
        row_cells = table.rows[1].cells
        row_cells[0].text = str(spatial_stats.get("ndvi_mean", "N/A"))
        row_cells[1].text = str(spatial_stats.get("ndwi_mean", "N/A"))
        row_cells[2].text = str(spatial_stats.get("ndbi_mean", "N/A"))
        row_cells[3].text = str(spatial_stats.get("nbr_mean", "N/A"))
        
        document.add_paragraph()
        
        INDEX_METADATA = {
            "NDVI": {
                "name": "Normalized Difference Vegetation Index (NDVI)",
                "desc": "Measures the density and health of vegetation in the area.",
                "legend": "🔴 Red/Orange: Barren land or built-up areas | 🟢 Green: Dense, healthy vegetation."
            },
            "NDWI": {
                "name": "Normalized Difference Water Index (NDWI)",
                "desc": "Highlights water bodies and tracks moisture content in plants and soil.",
                "legend": "⚪ White/Light Blue: Dry land or shallow water | 🔵 Deep Blue: Deep, clear water bodies."
            },
            "NDBI": {
                "name": "Normalized Difference Built-up Index (NDBI)",
                "desc": "Emphasizes urban areas, buildings, and impervious surfaces.",
                "legend": "⚪ White/Yellow: Non-urban areas or vegetation | 🔴 Red: Dense urban built-up areas."
            },
            "NBR": {
                "name": "Normalized Burn Ratio (NBR)",
                "desc": "Used to identify burned areas, fire severity, and vegetation survival.",
                "legend": "🟡 Yellow/Orange: Healthy vegetation or unburned land | ⚫ Dark Red/Black: Severe burn scars."
            }
        }
        
        if index_images:
            document.add_page_break()
            document.add_heading('GEE-Derived Scientific Results (Visualizations)', level=1)
            for idx_name, img_bytes in index_images.items():
                if img_bytes:
                    meta = INDEX_METADATA.get(idx_name, {"name": idx_name, "desc": "", "legend": ""})
                    document.add_heading(f'{meta["name"]}', level=2)
                    p_desc = document.add_paragraph(meta["desc"])
                    p_desc.style = 'List Bullet'
                    p_leg = document.add_paragraph(meta["legend"])
                    p_leg.style = 'List Bullet'
                    
                    try:
                        from PIL import Image
                        img_stream = io.BytesIO(img_bytes)
                        # Read the physical pixel dimensions to prevent Word from using weird DPI metadata
                        with Image.open(img_stream) as img:
                            px_width, px_height = img.size
                            aspect_ratio = px_height / px_width
                        
                        # Reset stream pointer
                        img_stream.seek(0)
                        
                        # Explicitly set width AND height so Word doesn't squash/stretch the image
                        target_width = 4.5
                        document.add_picture(img_stream, width=Inches(target_width), height=Inches(target_width * aspect_ratio))
                    except Exception as e:
                        logger.error(f"Failed to embed {idx_name} image: {e}")
                    document.add_paragraph()
                    
        docx_stream = io.BytesIO()
        document.save(docx_stream)
        docx_stream.seek(0)
        return docx_stream.read()

    @staticmethod
    def _add_parsed_text(paragraph, text: str):
        """Helper to parse **bold** text within a line and add to paragraph."""
        parts = re.split(r'(\*\*.*?\*\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                paragraph.add_run(part[2:-2]).bold = True
            else:
                paragraph.add_run(part)

    @staticmethod
    def generate_advanced_docx(package) -> bytes:
        """
        Generates a DOCX file deterministically from an assembled ReportPackage.
        """
        import httpx
        import asyncio
        from PIL import Image

        document = Document()
        
        # Title
        title = document.add_heading(package.title, 0)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        document.add_paragraph("Generated automatically by SatQuery Geospatial AI Pipeline.\n")
        
        # Metadata / Study Area
        meta_p = document.add_paragraph()
        meta_p.add_run("Job ID: ").bold = True
        meta_p.add_run(package.job_id + "\n")
        meta_p.add_run("Report Type: ").bold = True
        meta_p.add_run(package.report_type.capitalize() + "\n")
        meta_p.add_run("Period: ").bold = True
        meta_p.add_run(f"{package.analysis_period['start']} to {package.analysis_period['end']}\n")
        
        # Determine sections
        for section in package.sections:
            document.add_page_break()
            document.add_heading(section.title, level=1)
            
            if section.content:
                lines = section.content.split("\n")
                for line in lines:
                    line = line.strip()
                    if not line:
                        continue
                    if line.startswith("# "):
                        document.add_heading(line[2:].replace("**", ""), level=2)
                    elif line.startswith("## "):
                        document.add_heading(line[3:].replace("**", ""), level=3)
                    elif line.startswith("- ") or line.startswith("* "):
                        p = document.add_paragraph(style='List Bullet')
                        DocGenerator._add_parsed_text(p, line[2:])
                    else:
                        p = document.add_paragraph()
                        DocGenerator._add_parsed_text(p, line)
                        
            if section.metrics:
                table = document.add_table(rows=1, cols=3)
                table.style = 'Table Grid'
                hdr_cells = table.rows[0].cells
                hdr_cells[0].text = 'Metric'
                hdr_cells[1].text = 'Value'
                hdr_cells[2].text = 'Tool Source'
                
                for metric in section.metrics:
                    row_cells = table.add_row().cells
                    row_cells[0].text = str(metric.get("name", ""))
                    row_cells[1].text = str(metric.get("value", ""))
                    row_cells[2].text = str(metric.get("tool", ""))
                    
                document.add_paragraph()
                
            if section.findings:
                document.add_heading("Detailed Findings", level=2)
                for finding in section.findings:
                    p = document.add_paragraph(style='List Bullet')
                    DocGenerator._add_parsed_text(p, finding.get("text", ""))
                    if finding.get("evidence_ids"):
                        p_ev = document.add_paragraph()
                        p_ev.add_run("Supporting Evidence IDs: ").italic = True
                        p_ev.add_run(", ".join(finding["evidence_ids"])).italic = True
                        
            if section.visuals:
                for vis in section.visuals:
                    if vis.get("caption"):
                        document.add_heading(vis["caption"], level=2)
                        
                    url = vis.get("url")
                    if url:
                        try:
                            # Synchronous fetch for document assembly
                            import requests
                            resp = requests.get(url, timeout=15.0)
                            if resp.status_code == 200:
                                img_stream = io.BytesIO(resp.content)
                                with Image.open(img_stream) as img:
                                    px_width, px_height = img.size
                                    aspect_ratio = px_height / px_width
                                img_stream.seek(0)
                                target_width = 4.5
                                document.add_picture(img_stream, width=Inches(target_width), height=Inches(target_width * aspect_ratio))
                        except Exception as e:
                            logger.error(f"Failed to embed image from {url}: {e}")
                    document.add_paragraph()
                    
        # Provenance Section
        document.add_page_break()
        document.add_heading("Provenance & Audit Trail", level=1)
        p_prov = document.add_paragraph()
        p_prov.add_run("Provenance Status: ").bold = True
        p_prov.add_run(package.provenance.get("status", "unknown") + "\n")
        p_prov.add_run("Evidence Count: ").bold = True
        p_prov.add_run(str(package.provenance.get("evidence_count", 0)) + "\n")
        p_prov.add_run("Tools Executed: ").bold = True
        p_prov.add_run(", ".join(package.provenance.get("tools_used", [])) + "\n")
        
        docx_stream = io.BytesIO()
        document.save(docx_stream)
        docx_stream.seek(0)
        return docx_stream.read()
